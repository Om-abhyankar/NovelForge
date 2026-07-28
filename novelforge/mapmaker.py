"""
Fantasy map maker - data model, renderers and export.

The architecture that matters: a map is turned into a flat list of drawing
**primitives** by `build_primitives`. The Tkinter editor draws that list onto a
Canvas; the PNG exporter draws the same list with Pillow; the SVG exporter
writes the same list as XML. One display list, three backends, so what you see
while editing is exactly what you export.

Maps live in `13 Maps` as `<name>.map.json` alongside their exports. They are
deliberately kept out of project.json - a hand-drawn coastline is thousands of
points and would bloat the manifest that gets rewritten on every keystroke.
"""

from __future__ import annotations

import json
import math
import os
import random
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .atomic import read_json, safe_filename, write_json_atomic
from .model import from_dict, new_id, now_iso, to_dict

Point = Tuple[float, float]

# ==========================================================================
# Styles
# ==========================================================================

# Each style is a full palette.
#
# `paper` is the SEA - the page a map is drawn on reads as ocean once land is
# painted over it - so `terrain.land` must be clearly lighter than `paper` or
# the coastline reads as a bare outline with no land inside it. That contrast
# is the single most important thing separating a map from a diagram.
#
# `halo` is the concentric coastal glow, drawn under the land fill in
# progressively wider, lighter strokes. `terrain` overrides the per-kind
# defaults in TERRAIN so every style stays internally consistent.
STYLES: Dict[str, Dict[str, Any]] = {
    "parchment": {
        "label": "Parchment",
        "paper": "#d9c398",
        "paper_dark": "#bfa374",
        "ink": "#4a3520",
        "ink_light": "#7a5f3e",
        "water": "#a8c2ca",
        "water_deep": "#8fb0ba",
        "halo": ["#c9b48b", "#d2bf99", "#dbcaa8"],
        "grid": "#c4b18c",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#f2e8cb", "water": "#a8c2ca", "forest": "#b3c294",
            "desert": "#eadfae", "swamp": "#adb488", "ice": "#e4ebeb",
        },
    },
    "ink": {
        "label": "Ink on white",
        "paper": "#e9eef1",
        "paper_dark": "#d8e0e4",
        "ink": "#1e1e1e",
        "ink_light": "#5a5a5a",
        "water": "#d3e0e6",
        "water_deep": "#bfd2da",
        "halo": ["#dce5e9", "#e3eaee", "#eaeff2"],
        "grid": "#cbd4d8",
        "texture": False,
        "vignette": False,
        "terrain": {
            "land": "#fdfdfb", "water": "#d3e0e6", "forest": "#dde8d5",
            "desert": "#f5efdc", "swamp": "#dfe2cf", "ice": "#f2f7f8",
        },
    },
    "dark": {
        "label": "Dark atlas",
        "paper": "#161c22",
        "paper_dark": "#101519",
        "ink": "#ddd6c6",
        "ink_light": "#8f887c",
        "water": "#1b2731",
        "water_deep": "#141d25",
        "halo": ["#222c35", "#1d262e", "#192128"],
        "grid": "#2b343c",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#333c44", "water": "#1b2731", "forest": "#2f4038",
            "desert": "#45412f", "swamp": "#333a2e", "ice": "#3f4a52",
        },
    },
    "treasure": {
        "label": "Treasure map",
        "paper": "#c9a86f",
        "paper_dark": "#a88a52",
        "ink": "#5a3a18",
        "ink_light": "#8a6535",
        "water": "#b9a06d",
        "water_deep": "#a48c5c",
        "halo": ["#bd9c63", "#c6a771", "#cfb281"],
        "grid": "#b5975f",
        "texture": True,
        "vignette": True,
        "terrain": {
            "land": "#e8d3a4", "water": "#b9a06d", "forest": "#b0ac78",
            "desert": "#e2cd97", "swamp": "#a5a274", "ice": "#dcdcc9",
        },
    },
}

TERRAIN: Dict[str, Dict[str, Any]] = {
    "land":      {"label": "Land / coast", "fill": "#f2e8cb", "closed": True,
                  "halo": True, "decor": None},
    "water":     {"label": "Sea / lake", "fill": "#a8c2ca", "closed": True,
                  "halo": False, "decor": None},
    "forest":    {"label": "Forest", "fill": "#b3c294", "closed": True,
                  "halo": False, "decor": "trees"},
    "mountains": {"label": "Mountain range", "fill": None, "closed": False,
                  "halo": False, "decor": "peaks"},
    "hills":     {"label": "Hills", "fill": None, "closed": False,
                  "halo": False, "decor": "hills"},
    "desert":    {"label": "Desert", "fill": "#e4d5a8", "closed": True,
                  "halo": False, "decor": "dots"},
    "swamp":     {"label": "Marsh / swamp", "fill": "#b3b585", "closed": True,
                  "halo": False, "decor": "squiggles"},
    "ice":       {"label": "Ice / tundra", "fill": "#dfe7e8", "closed": True,
                  "halo": False, "decor": None},
    "region":    {"label": "Region / border", "fill": None, "closed": True,
                  "halo": False, "decor": None},
    "river":     {"label": "River", "fill": None, "closed": False,
                  "halo": False, "decor": "river"},
    "road":      {"label": "Road", "fill": None, "closed": False,
                  "halo": False, "decor": "road"},
    "wall":      {"label": "Wall", "fill": None, "closed": False,
                  "halo": False, "decor": "wall"},
    "route":     {"label": "Route / journey", "fill": None, "closed": False,
                  "halo": False, "decor": "route"},
}

TERRAIN_ORDER = [
    "land", "water", "forest", "mountains", "hills", "desert", "swamp",
    "ice", "region", "river", "road", "wall", "route",
]

# Draw order. Water under land, decorations over both, routes on top.
PAINT_ORDER = {
    "water": 0, "land": 1, "ice": 2, "desert": 3, "swamp": 4, "forest": 5,
    "hills": 6, "mountains": 7, "region": 8, "river": 9, "road": 10,
    "wall": 11, "route": 12,
}

PIN_KINDS: Dict[str, str] = {
    "capital": "Capital city",
    "city": "City",
    "town": "Town",
    "village": "Village",
    "castle": "Castle / keep",
    "tower": "Tower",
    "temple": "Temple",
    "port": "Port / harbour",
    "ruin": "Ruin",
    "cave": "Cave",
    "dungeon": "Dungeon",
    "mine": "Mine",
    "camp": "Camp",
    "bridge": "Bridge",
    "inn": "Inn",
    "battle": "Battle site",
    "danger": "Danger",
    "treasure": "Treasure",
    "landmark": "Landmark",
    "portal": "Portal",
}

MAP_KINDS = {
    "world": "World",
    "continent": "Continent",
    "region": "Region",
    "city": "City / town plan",
    "building": "Building / interior",
    "dungeon": "Dungeon",
    "treasure": "Treasure map",
    "battle": "Battle plan",
}

GRID_KINDS = {"none": "None", "square": "Square", "hex": "Hex"}


# ==========================================================================
# Model
# ==========================================================================


@dataclass
class Layer:
    name: str = "Base"
    visible: bool = True
    locked: bool = False


@dataclass
class Shape:
    id: str = ""
    kind: str = "land"
    points: List[Point] = field(default_factory=list)
    closed: bool = True
    fill: str = ""          # blank means "use the terrain default"
    outline: str = ""
    width: float = 2.0
    label: str = ""
    layer: str = "Base"

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("shp")
        # JSON round-trips tuples into lists; normalise so geometry maths works.
        self.points = [(float(p[0]), float(p[1])) for p in self.points if len(p) >= 2]

    def bounds(self) -> Tuple[float, float, float, float]:
        if not self.points:
            return (0.0, 0.0, 0.0, 0.0)
        xs = [p[0] for p in self.points]
        ys = [p[1] for p in self.points]
        return (min(xs), min(ys), max(xs), max(ys))

    def centroid(self) -> Point:
        if not self.points:
            return (0.0, 0.0)
        return (sum(p[0] for p in self.points) / len(self.points),
                sum(p[1] for p in self.points) / len(self.points))


@dataclass
class Pin:
    id: str = ""
    x: float = 0.0
    y: float = 0.0
    kind: str = "city"
    label: str = ""
    label_side: str = "e"      # e | w | n | s
    entity_id: str = ""        # links to a Location entity in the project
    notes: str = ""
    size: float = 7.0
    layer: str = "Base"

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("pin")


@dataclass
class MapLabel:
    id: str = ""
    x: float = 0.0
    y: float = 0.0
    text: str = ""
    size: int = 16
    color: str = ""
    italic: bool = False
    bold: bool = False
    tracking: float = 0.0      # extra letter spacing, for region names
    layer: str = "Base"

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("lbl")


@dataclass
class GameMap:
    id: str = ""
    name: str = "New Map"
    kind: str = "world"
    style: str = "parchment"
    width: int = 1600
    height: int = 1100
    grid: str = "none"
    grid_size: int = 80
    scale_text: str = ""
    compass: bool = True
    border: bool = True
    title_on_map: bool = True
    # When a place name has nowhere to go, hide it rather than print it on top
    # of another. Turn off to see every name, overlaps and all.
    hide_colliding_labels: bool = True
    auto_place_labels: bool = True
    notes: str = ""
    seed: int = 7
    layers: List[Layer] = field(default_factory=lambda: [Layer()])
    shapes: List[Shape] = field(default_factory=list)
    pins: List[Pin] = field(default_factory=list)
    labels: List[MapLabel] = field(default_factory=list)
    created: str = field(default_factory=now_iso)
    modified: str = field(default_factory=now_iso)

    def __post_init__(self) -> None:
        if not self.id:
            self.id = new_id("map")
        if not self.layers:
            self.layers = [Layer()]

    # -- serialisation --------------------------------------------------
    def to_json(self) -> Dict[str, Any]:
        return {
            "id": self.id, "name": self.name, "kind": self.kind,
            "style": self.style, "width": self.width, "height": self.height,
            "grid": self.grid, "grid_size": self.grid_size,
            "scale_text": self.scale_text, "compass": self.compass,
            "border": self.border, "title_on_map": self.title_on_map,
            "hide_colliding_labels": self.hide_colliding_labels,
            "auto_place_labels": self.auto_place_labels,
            "notes": self.notes, "seed": self.seed,
            "layers": [to_dict(l) for l in self.layers],
            "shapes": [
                {**to_dict(s), "points": [[p[0], p[1]] for p in s.points]}
                for s in self.shapes
            ],
            "pins": [to_dict(p) for p in self.pins],
            "labels": [to_dict(l) for l in self.labels],
            "created": self.created, "modified": self.modified,
        }

    @classmethod
    def from_json(cls, data: Dict[str, Any]) -> "GameMap":
        skip = {"layers", "shapes", "pins", "labels"}
        obj = cls(**{
            k: v for k, v in (data or {}).items()
            if k not in skip and k in {f.name for f in cls.__dataclass_fields__.values()}
        })
        obj.layers = [from_dict(Layer, d) for d in (data.get("layers") or [])] \
            or [Layer()]
        obj.shapes = [from_dict(Shape, d) for d in (data.get("shapes") or [])]
        obj.pins = [from_dict(Pin, d) for d in (data.get("pins") or [])]
        obj.labels = [from_dict(MapLabel, d) for d in (data.get("labels") or [])]
        return obj

    # -- helpers --------------------------------------------------------
    def palette(self) -> Dict[str, Any]:
        return STYLES.get(self.style, STYLES["parchment"])

    def layer_names(self) -> List[str]:
        return [l.name for l in self.layers]

    def visible_layers(self) -> set:
        return {l.name for l in self.layers if l.visible}

    def layer(self, name: str) -> Optional[Layer]:
        return next((l for l in self.layers if l.name == name), None)

    def shape(self, shape_id: str) -> Optional[Shape]:
        return next((s for s in self.shapes if s.id == shape_id), None)

    def pin(self, pin_id: str) -> Optional[Pin]:
        return next((p for p in self.pins if p.id == pin_id), None)

    def label(self, label_id: str) -> Optional[MapLabel]:
        return next((l for l in self.labels if l.id == label_id), None)

    def touch(self) -> None:
        self.modified = now_iso()

    def is_empty(self) -> bool:
        return not (self.shapes or self.pins or self.labels)


# ==========================================================================
# Primitives - the shared display list
# ==========================================================================

# ("polygon", points, fill, outline, width, dash)
# ("line",    points, color, width, dash)
# ("ellipse", x0, y0, x1, y1, fill, outline, width)
# ("text",    x, y, text, size, color, anchor, italic, bold, tracking)


def _rgb(colour: str) -> Tuple[int, int, int]:
    colour = (colour or "#000000").lstrip("#")
    if len(colour) == 3:
        colour = "".join(c * 2 for c in colour)
    if len(colour) != 6:
        return (0, 0, 0)
    try:
        return (int(colour[0:2], 16), int(colour[2:4], 16), int(colour[4:6], 16))
    except ValueError:
        return (0, 0, 0)


def _mix(a: str, b: str, t: float) -> str:
    ra, ga, ba = _rgb(a)
    rb, gb, bb = _rgb(b)
    return "#%02x%02x%02x" % (
        int(ra + (rb - ra) * t), int(ga + (gb - ga) * t), int(ba + (bb - ba) * t)
    )


def _smooth(points: Sequence[Point], iterations: int = 2) -> List[Point]:
    """
    Chaikin corner-cutting, so a hand-drawn coastline is not visibly polygonal.

    Bounded by `iterations` and skipped for very long paths, because each pass
    doubles the point count and a freehand coast can already hold thousands.
    """
    pts = [(float(p[0]), float(p[1])) for p in points]
    if len(pts) < 3 or len(pts) > 600:
        return pts
    for _ in range(max(0, min(3, iterations))):
        out: List[Point] = [pts[0]]
        for a, b in zip(pts, pts[1:]):
            out.append((a[0] * 0.75 + b[0] * 0.25, a[1] * 0.75 + b[1] * 0.25))
            out.append((a[0] * 0.25 + b[0] * 0.75, a[1] * 0.25 + b[1] * 0.75))
        out.append(pts[-1])
        pts = out
        if len(pts) > 2400:
            break
    return pts


def _path_length(points: Sequence[Point]) -> float:
    return sum(
        math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(points, points[1:])
    )


def _walk_path(points: Sequence[Point], spacing: float) -> List[Tuple[Point, float]]:
    """
    Sample points along a polyline at roughly `spacing` apart.

    Returns (point, angle). Used to scatter mountain peaks and trees along a
    range. `spacing` is clamped so a degenerate value cannot loop forever.
    """
    spacing = max(4.0, float(spacing))
    out: List[Tuple[Point, float]] = []
    carry = 0.0
    for a, b in zip(points, points[1:]):
        dx, dy = b[0] - a[0], b[1] - a[1]
        seg = math.hypot(dx, dy)
        if seg < 1e-6:
            continue
        angle = math.atan2(dy, dx)
        position = carry
        while position < seg:
            t = position / seg
            out.append(((a[0] + dx * t, a[1] + dy * t), angle))
            position += spacing
            if len(out) > 4000:      # hard cap: decoration, not geometry
                return out
        carry = position - seg
    return out


def _scatter_in_bounds(shape: Shape, spacing: float, seed: int
                       ) -> List[Point]:
    """Jittered grid of points inside a shape's bounding box, filtered by the polygon."""
    x0, y0, x1, y1 = shape.bounds()
    if x1 - x0 < 2 or y1 - y0 < 2:
        return []
    rng = random.Random(seed ^ (hash(shape.id) & 0xFFFF))
    spacing = max(10.0, float(spacing))
    out: List[Point] = []
    rows = int((y1 - y0) / spacing) + 1
    cols = int((x1 - x0) / spacing) + 1
    if rows * cols > 6000:           # keep decoration bounded on huge shapes
        spacing = math.sqrt((x1 - x0) * (y1 - y0) / 6000.0)
        rows = int((y1 - y0) / spacing) + 1
        cols = int((x1 - x0) / spacing) + 1
    for r in range(rows):
        for c in range(cols):
            px = x0 + c * spacing + rng.uniform(-spacing * 0.3, spacing * 0.3)
            py = y0 + r * spacing + rng.uniform(-spacing * 0.3, spacing * 0.3)
            if point_in_polygon((px, py), shape.points):
                out.append((px, py))
    return out


def point_in_polygon(pt: Point, polygon: Sequence[Point]) -> bool:
    """Ray casting. Used for hit-testing and for filling terrain decoration."""
    if len(polygon) < 3:
        return False
    x, y = pt
    inside = False
    n = len(polygon)
    j = n - 1
    for i in range(n):
        xi, yi = polygon[i]
        xj, yj = polygon[j]
        if ((yi > y) != (yj > y)) and \
                (x < (xj - xi) * (y - yi) / ((yj - yi) or 1e-12) + xi):
            inside = not inside
        j = i
    return inside


def distance_to_path(pt: Point, points: Sequence[Point]) -> float:
    """Shortest distance from a point to a polyline, for hit-testing lines."""
    if not points:
        return float("inf")
    if len(points) == 1:
        return math.hypot(pt[0] - points[0][0], pt[1] - points[0][1])
    best = float("inf")
    px, py = pt
    for a, b in zip(points, points[1:]):
        ax, ay = a
        bx, by = b
        dx, dy = bx - ax, by - ay
        seg_sq = dx * dx + dy * dy
        if seg_sq < 1e-12:
            best = min(best, math.hypot(px - ax, py - ay))
            continue
        t = max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / seg_sq))
        best = min(best, math.hypot(px - (ax + t * dx), py - (ay + t * dy)))
    return best


# -- pin glyphs ------------------------------------------------------------


def pin_primitives(pin: Pin, ink: str, paper: str,
                   accent: str = "") -> List[tuple]:
    """
    Build the drawing primitives for one map pin.

    Shared by the editor and both exporters, so a treasure X on screen is the
    same treasure X in the PNG.
    """
    x, y = float(pin.x), float(pin.y)
    s = max(3.0, float(pin.size))
    kind = pin.kind
    out: List[tuple] = []
    accent = accent or ink

    def ring(radius: float, fill: str, width: float = 1.6) -> None:
        out.append(("ellipse", x - radius, y - radius, x + radius, y + radius,
                    fill, ink, width))

    if kind == "capital":
        ring(s * 1.55, paper, 1.8)
        ring(s * 0.75, ink, 1.2)
        for i in range(8):                      # star burst
            a = i * math.pi / 4
            out.append(("line",
                        [(x + math.cos(a) * s * 1.6, y + math.sin(a) * s * 1.6),
                         (x + math.cos(a) * s * 2.15, y + math.sin(a) * s * 2.15)],
                        ink, 1.4, False))
    elif kind == "city":
        ring(s * 1.15, paper, 1.8)
        ring(s * 0.45, ink, 1.0)
    elif kind == "town":
        ring(s * 0.8, paper, 1.6)
        out.append(("ellipse", x - s * 0.22, y - s * 0.22,
                    x + s * 0.22, y + s * 0.22, ink, ink, 1.0))
    elif kind == "village":
        out.append(("polygon",
                    [(x - s * 0.6, y + s * 0.5), (x + s * 0.6, y + s * 0.5),
                     (x + s * 0.6, y - s * 0.15), (x, y - s * 0.75),
                     (x - s * 0.6, y - s * 0.15)],
                    paper, ink, 1.4, False))
    elif kind == "castle":
        base_y = y + s * 0.7
        top_y = y - s * 0.35
        out.append(("polygon",
                    [(x - s, base_y), (x - s, top_y), (x - s * 0.6, top_y),
                     (x - s * 0.6, top_y - s * 0.4), (x - s * 0.2, top_y - s * 0.4),
                     (x - s * 0.2, top_y), (x + s * 0.2, top_y),
                     (x + s * 0.2, top_y - s * 0.4), (x + s * 0.6, top_y - s * 0.4),
                     (x + s * 0.6, top_y), (x + s, top_y), (x + s, base_y)],
                    paper, ink, 1.5, False))
    elif kind == "tower":
        out.append(("polygon",
                    [(x - s * 0.45, y + s * 0.9), (x - s * 0.45, y - s * 0.5),
                     (x, y - s * 1.15), (x + s * 0.45, y - s * 0.5),
                     (x + s * 0.45, y + s * 0.9)],
                    paper, ink, 1.5, False))
    elif kind == "temple":
        out.append(("polygon",
                    [(x - s, y + s * 0.3), (x, y - s * 0.95), (x + s, y + s * 0.3)],
                    paper, ink, 1.5, False))
        out.append(("line", [(x - s, y + s * 0.7), (x + s, y + s * 0.7)],
                    ink, 1.6, False))
    elif kind == "port":
        out.append(("line", [(x, y - s), (x, y + s * 0.9)], ink, 1.8, False))
        out.append(("line", [(x - s * 0.7, y - s * 0.55),
                             (x + s * 0.7, y - s * 0.55)], ink, 1.6, False))
        out.append(("line", [(x - s * 0.85, y + s * 0.25),
                             (x, y + s * 0.95), (x + s * 0.85, y + s * 0.25)],
                    ink, 1.7, False))
    elif kind == "ruin":
        out.append(("line", [(x - s * 0.8, y + s * 0.7),
                             (x - s * 0.8, y - s * 0.5)], ink, 1.8, False))
        out.append(("line", [(x - s * 0.1, y + s * 0.7),
                             (x - s * 0.1, y - s * 0.1)], ink, 1.8, False))
        out.append(("line", [(x + s * 0.7, y + s * 0.7),
                             (x + s * 0.7, y - s * 0.65)], ink, 1.8, False))
        out.append(("line", [(x - s, y + s * 0.8), (x + s, y + s * 0.8)],
                    ink, 1.5, False))
    elif kind in ("cave", "dungeon"):
        out.append(("polygon",
                    [(x - s, y + s * 0.75), (x - s * 0.75, y - s * 0.35),
                     (x, y - s * 0.8), (x + s * 0.75, y - s * 0.35),
                     (x + s, y + s * 0.75)],
                    ink if kind == "dungeon" else paper, ink, 1.5, False))
        if kind == "dungeon":
            out.append(("line", [(x - s * 0.35, y + s * 0.75),
                                 (x - s * 0.35, y + s * 0.1)], paper, 1.4, False))
            out.append(("line", [(x + s * 0.35, y + s * 0.75),
                                 (x + s * 0.35, y + s * 0.1)], paper, 1.4, False))
    elif kind == "mine":
        out.append(("polygon",
                    [(x - s, y + s * 0.7), (x, y - s * 0.8), (x + s, y + s * 0.7)],
                    paper, ink, 1.5, False))
        out.append(("line", [(x - s * 0.4, y + s * 0.7),
                             (x + s * 0.4, y - s * 0.1)], ink, 1.5, False))
    elif kind == "camp":
        out.append(("polygon",
                    [(x - s * 0.9, y + s * 0.7), (x, y - s * 0.85),
                     (x + s * 0.9, y + s * 0.7)],
                    "", ink, 1.6, False))
        out.append(("line", [(x, y - s * 0.85), (x, y + s * 0.7)], ink, 1.2, False))
    elif kind == "bridge":
        out.append(("line", [(x - s, y + s * 0.4), (x - s * 0.4, y - s * 0.5),
                             (x + s * 0.4, y - s * 0.5), (x + s, y + s * 0.4)],
                    ink, 1.8, False))
    elif kind == "inn":
        ring(s * 0.85, paper, 1.5)
        out.append(("line", [(x - s * 0.4, y - s * 0.3),
                             (x - s * 0.4, y + s * 0.35)], ink, 1.3, False))
        out.append(("line", [(x + s * 0.4, y - s * 0.3),
                             (x + s * 0.4, y + s * 0.35)], ink, 1.3, False))
        out.append(("line", [(x - s * 0.4, y + s * 0.05),
                             (x + s * 0.4, y + s * 0.05)], ink, 1.3, False))
    elif kind == "battle":
        out.append(("line", [(x - s, y - s), (x + s, y + s)], ink, 2.0, False))
        out.append(("line", [(x + s, y - s), (x - s, y + s)], ink, 2.0, False))
        out.append(("line", [(x - s * 1.15, y - s * 0.55),
                             (x - s * 0.55, y - s * 1.15)], ink, 1.5, False))
    elif kind == "treasure":
        out.append(("line", [(x - s * 1.1, y - s * 1.1),
                             (x + s * 1.1, y + s * 1.1)], accent, 3.0, False))
        out.append(("line", [(x + s * 1.1, y - s * 1.1),
                             (x - s * 1.1, y + s * 1.1)], accent, 3.0, False))
    elif kind == "danger":
        out.append(("polygon",
                    [(x, y - s * 1.1), (x + s, y + s * 0.75), (x - s, y + s * 0.75)],
                    paper, ink, 1.7, False))
        out.append(("line", [(x, y - s * 0.45), (x, y + s * 0.2)], ink, 1.8, False))
        out.append(("ellipse", x - s * 0.13, y + s * 0.38,
                    x + s * 0.13, y + s * 0.58, ink, ink, 1.0))
    elif kind == "portal":
        ring(s * 1.1, "", 1.8)
        ring(s * 0.62, "", 1.3)
        ring(s * 0.2, ink, 1.0)
    else:  # landmark
        out.append(("polygon",
                    [(x, y - s), (x + s * 0.32, y - s * 0.32),
                     (x + s, y), (x + s * 0.32, y + s * 0.32),
                     (x, y + s), (x - s * 0.32, y + s * 0.32),
                     (x - s, y), (x - s * 0.32, y - s * 0.32)],
                    paper, ink, 1.4, False))
    return out


def pin_label_anchor(pin: Pin, side: str = "") -> Tuple[float, float, str]:
    """Where a pin's label sits, and its text anchor."""
    s = max(3.0, float(pin.size))
    gap = s * 2.0
    side = (side or pin.label_side or "e").lower()
    if side == "w":
        return (pin.x - gap, pin.y, "e")
    if side == "n":
        return (pin.x, pin.y - gap, "s")
    if side == "s":
        return (pin.x, pin.y + gap, "n")
    if side == "ne":
        return (pin.x + gap * 0.72, pin.y - gap * 0.72, "w")
    if side == "nw":
        return (pin.x - gap * 0.72, pin.y - gap * 0.72, "e")
    if side == "se":
        return (pin.x + gap * 0.72, pin.y + gap * 0.72, "w")
    if side == "sw":
        return (pin.x - gap * 0.72, pin.y + gap * 0.72, "e")
    return (pin.x + gap, pin.y, "w")


# --------------------------------------------------------------------------
# Automatic label placement
#
# A generated map puts fifty place names on the page and they collide. Real
# cartographers solve this by moving a label around its point and, failing
# that, leaving it off. That is exactly what this does.
# --------------------------------------------------------------------------

Box = Tuple[float, float, float, float]

# Which labels win a fight. A capital's name matters more than a village's.
LABEL_PRIORITY = {
    "capital": 0, "city": 1, "port": 2, "castle": 3, "temple": 4,
    "town": 5, "dungeon": 6, "ruin": 7, "portal": 8, "landmark": 9,
    "mine": 10, "battle": 11, "treasure": 12, "bridge": 13, "cave": 14,
    "tower": 15, "inn": 16, "camp": 17, "danger": 18, "village": 19,
}

# Order in which alternative positions are tried around a point.
LABEL_SIDES = ("e", "w", "n", "s", "ne", "nw", "se", "sw")


def text_box(x: float, y: float, text: str, size: float, anchor: str,
             tracking: float = 0.0) -> Box:
    """
    Estimate a text bounding box without a font.

    0.52em average advance width is a good approximation for a serif face and
    keeps this dependency-free, which matters because the same estimate has to
    hold for the Tk canvas, Pillow and SVG.
    """
    if not text:
        return (x, y, x, y)
    width = len(text) * size * 0.52 + tracking * max(0, len(text) - 1)
    height = size * 1.15
    if anchor == "w":
        x0 = x
    elif anchor == "e":
        x0 = x - width
    else:
        x0 = x - width / 2.0
    if anchor == "n":
        y0 = y
    elif anchor == "s":
        y0 = y - height
    else:
        y0 = y - height / 2.0
    return (x0, y0, x0 + width, y0 + height)


def boxes_overlap(a: Box, b: Box, pad: float = 1.5) -> bool:
    return not (a[2] + pad < b[0] or b[2] + pad < a[0]
                or a[3] + pad < b[1] or b[3] + pad < a[1])


def layout_pin_labels(gm: GameMap, visible: Optional[set] = None
                      ) -> Dict[str, Tuple[str, bool]]:
    """
    Choose a side for every pin label and decide which ones must be dropped.

    Returns {pin_id: (side, show)}. Reserved first, in order: the map title,
    compass and scale bar; then free-standing text labels; then every pin's own
    glyph, so a name never sits on top of a marker.
    """
    if visible is None:
        visible = gm.visible_layers()
    placed: List[Box] = []

    # Furniture.
    if gm.title_on_map and gm.name.strip():
        size = max(18, int(gm.height * 0.038))
        placed.append(text_box(gm.width / 2.0, gm.height * 0.062,
                               gm.name.upper(), size, "center", size * 0.14))
    if gm.compass:
        r = min(gm.width, gm.height) * 0.055
        cx, cy = gm.width - r * 2.1, gm.height - r * 2.1
        placed.append((cx - r * 1.5, cy - r * 1.8, cx + r * 1.5, cy + r * 1.5))
    if gm.scale_text.strip():
        bar = min(gm.width * 0.16, 230.0)
        x0 = gm.width * 0.045
        y0 = gm.height - gm.height * 0.052
        placed.append((x0 - 6, y0 - gm.height * 0.03, x0 + bar + 6,
                       y0 + gm.height * 0.02))

    # Free-standing labels (region and ocean names) always win.
    for label in gm.labels:
        if label.layer not in visible or not label.text.strip():
            continue
        placed.append(text_box(label.x, label.y, label.text, label.size,
                               "center", label.tracking))

    # Named areas drawn from a shape's label.
    for shape in gm.shapes:
        if shape.layer in visible and shape.label.strip():
            cx, cy = shape.centroid()
            placed.append(text_box(cx, cy, shape.label,
                                   max(11, int(gm.height * 0.019)),
                                   "center", 2.0))

    pins = [p for p in gm.pins if p.layer in visible]
    # Reserve every glyph before placing any text.
    for pin in pins:
        s = max(3.0, float(pin.size)) * 1.7
        placed.append((pin.x - s, pin.y - s, pin.x + s, pin.y + s))

    out: Dict[str, Tuple[str, bool]] = {}
    ordered = sorted(
        pins,
        key=lambda p: (LABEL_PRIORITY.get(p.kind, 50), -float(p.size),
                       p.label.lower()),
    )
    size = max(9, int(gm.height * 0.0155))
    for pin in ordered:
        if not pin.label.strip():
            out[pin.id] = (pin.label_side or "e", False)
            continue
        # Try the author's chosen side first, then the rest.
        preferred = (pin.label_side or "e").lower()
        candidates = [preferred] + [s for s in LABEL_SIDES if s != preferred]
        chosen: Optional[str] = None
        for side in candidates:
            lx, ly, anchor = pin_label_anchor(pin, side)
            box = text_box(lx, ly, pin.label, size, anchor)
            # Stay on the page.
            if box[0] < 4 or box[1] < 4 or box[2] > gm.width - 4 \
                    or box[3] > gm.height - 4:
                continue
            if any(boxes_overlap(box, other) for other in placed):
                continue
            chosen = side
            placed.append(box)
            break
        if chosen is None:
            # Nowhere to put it. Hiding one village name is far better than a
            # page of names printed on top of each other.
            out[pin.id] = (preferred, not gm.hide_colliding_labels)
            if not gm.hide_colliding_labels:
                lx, ly, anchor = pin_label_anchor(pin, preferred)
                placed.append(text_box(lx, ly, pin.label, size, anchor))
        else:
            out[pin.id] = (chosen, True)
    return out


# -- terrain decoration ----------------------------------------------------


def _peaks(points: Sequence[Point], ink: str, paper: str,
           spacing: float = 24.0, height: float = 15.0,
           seed: int = 7) -> List[tuple]:
    """
    Mountains drawn along a ridge line.

    Sizes and vertical offsets vary per peak, and peaks are drawn back-to-front
    so the nearer ones overlap the farther ones. A row of identical triangles
    reads as a fence; varied overlapping ones read as a range.
    """
    rng = random.Random(seed)
    samples = _walk_path(points, spacing)
    drawn: List[Tuple[float, float, float]] = []
    for (px, py), _angle in samples:
        scale = rng.uniform(0.62, 1.42)
        drawn.append((px, py + rng.uniform(-height * 0.28, height * 0.28),
                      height * scale))
    # Farther (higher on the page) first, so nearer peaks sit in front.
    drawn.sort(key=lambda item: item[1])

    out: List[tuple] = []
    for px, py, size in drawn:
        half = size * 0.72
        out.append(("polygon",
                    [(px - half, py + size * 0.42), (px, py - size * 0.62),
                     (px + half, py + size * 0.42)],
                    paper, ink, 1.4, False))
        # Shade the right flank so the range has a light direction.
        out.append(("polygon",
                    [(px, py - size * 0.62), (px + half, py + size * 0.42),
                     (px + half * 0.28, py + size * 0.42)],
                    _mix(paper, ink, 0.22), "", 0.0, False))
        out.append(("line", [(px - half * 0.26, py + size * 0.08),
                             (px, py - size * 0.62)], ink, 0.9, False))
    return out


def _hills(points: Sequence[Point], ink: str,
           spacing: float = 22.0) -> List[tuple]:
    out: List[tuple] = []
    for (px, py), _a in _walk_path(points, spacing):
        r = 7.0
        out.append(("line",
                    [(px - r, py + 2), (px - r * 0.45, py - r * 0.55),
                     (px + r * 0.1, py + 2)], ink, 1.5, False))
        out.append(("line",
                    [(px + r * 0.05, py + 2), (px + r * 0.6, py - r * 0.35),
                     (px + r * 1.15, py + 2)], ink, 1.3, False))
    return out


def _trees(shape: Shape, ink: str, seed: int,
           spacing: float = 30.0) -> List[tuple]:
    """
    Trees scattered inside a forest.

    Clumped rather than evenly spread: a coarse hash gates whole neighbourhoods
    on or off, so the wood has clearings and dense stands instead of looking
    like tiled wallpaper. Sizes vary and about a fifth of candidate spots are
    skipped outright.
    """
    rng = random.Random(seed ^ (hash(shape.id) & 0xFFFF))
    out: List[tuple] = []
    for px, py in _scatter_in_bounds(shape, spacing, seed):
        # Clumping: a smooth-ish pseudo-random field over a ~4-cell grid.
        clump = (math.sin(px * 0.021 + seed * 0.7) +
                 math.cos(py * 0.019 - seed * 0.4))
        if clump < -0.55:
            continue
        if rng.random() < 0.18:
            continue
        r = 3.3 * rng.uniform(0.78, 1.5)
        jx = px + rng.uniform(-spacing * 0.22, spacing * 0.22)
        jy = py + rng.uniform(-spacing * 0.22, spacing * 0.22)
        if rng.random() < 0.3:
            # A conifer, for variety.
            out.append(("polygon",
                        [(jx - r * 0.8, jy + r * 0.7), (jx, jy - r * 1.5),
                         (jx + r * 0.8, jy + r * 0.7)], "", ink, 1.1, False))
        else:
            out.append(("ellipse", jx - r, jy - r * 1.25, jx + r, jy + r * 0.55,
                        "", ink, 1.1))
        out.append(("line", [(jx, jy + r * 0.6), (jx, jy + r * 1.25)],
                    ink, 1.0, False))
    return out


def _dots(shape: Shape, ink: str, seed: int,
          spacing: float = 20.0) -> List[tuple]:
    out: List[tuple] = []
    for px, py in _scatter_in_bounds(shape, spacing, seed + 31):
        out.append(("ellipse", px - 1.0, py - 1.0, px + 1.0, py + 1.0,
                    ink, ink, 1.0))
    return out


def _squiggles(shape: Shape, ink: str, seed: int,
               spacing: float = 26.0) -> List[tuple]:
    out: List[tuple] = []
    for px, py in _scatter_in_bounds(shape, spacing, seed + 71):
        out.append(("line", [(px - 6, py), (px - 2, py - 2.5), (px + 2, py),
                             (px + 6, py - 2.5)], ink, 1.2, False))
    return out


def _tapered_river(points: Sequence[Point], colour: str,
                   width: float) -> List[tuple]:
    """A river drawn in three passes so it thickens toward its mouth."""
    pts = list(points)
    if len(pts) < 2:
        return []
    out: List[tuple] = []
    third = max(2, len(pts) // 3)
    out.append(("line", pts, colour, max(1.0, width * 0.6), False))
    out.append(("line", pts[third:], colour, max(1.2, width * 0.85), False))
    out.append(("line", pts[third * 2:], colour, max(1.5, width * 1.15), False))
    return out


# -- background and furniture ----------------------------------------------


def _grid_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    # Blended toward the paper: a grid is a reading aid, not a feature, and at
    # full strength it competes with the terrain for attention.
    colour = _mix(palette["grid"], palette["paper"], 0.55)
    step = max(20, int(gm.grid_size))
    out: List[tuple] = []
    if gm.grid == "square":
        x = 0
        while x <= gm.width:
            out.append(("line", [(x, 0), (x, gm.height)], colour, 0.7, False))
            x += step
        y = 0
        while y <= gm.height:
            out.append(("line", [(0, y), (gm.width, y)], colour, 0.7, False))
            y += step
    elif gm.grid == "hex":
        r = step / 2.0
        h = r * math.sqrt(3)
        row = 0
        y = 0.0
        while y <= gm.height + h:
            offset = 0.0 if row % 2 == 0 else r * 1.5
            x = offset
            while x <= gm.width + r * 2:
                pts = [
                    (x + r * math.cos(math.pi / 3 * i),
                     y + r * math.sin(math.pi / 3 * i))
                    for i in range(6)
                ]
                out.append(("polygon", pts, "", colour, 0.7, False))
                x += r * 3
            y += h / 2
            row += 1
            if row > 400:            # guard against a pathological grid_size
                break
    return out


def _compass_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    ink, paper = palette["ink"], palette["paper"]
    r = min(gm.width, gm.height) * 0.055
    cx = gm.width - r * 2.1
    cy = gm.height - r * 2.1
    out: List[tuple] = [
        ("ellipse", cx - r, cy - r, cx + r, cy + r, "", ink, 1.4),
        ("ellipse", cx - r * 0.75, cy - r * 0.75, cx + r * 0.75, cy + r * 0.75,
         "", palette["ink_light"], 0.9),
    ]
    for i in range(4):
        a = -math.pi / 2 + i * math.pi / 2
        tip = (cx + math.cos(a) * r * 0.95, cy + math.sin(a) * r * 0.95)
        left = (cx + math.cos(a + math.pi / 2) * r * 0.17,
                cy + math.sin(a + math.pi / 2) * r * 0.17)
        right = (cx + math.cos(a - math.pi / 2) * r * 0.17,
                 cy + math.sin(a - math.pi / 2) * r * 0.17)
        out.append(("polygon", [tip, left, (cx, cy), right],
                    ink if i == 0 else paper, ink, 1.1, False))
    for i in range(4):
        a = -math.pi / 4 + i * math.pi / 2
        out.append(("line",
                    [(cx + math.cos(a) * r * 0.2, cy + math.sin(a) * r * 0.2),
                     (cx + math.cos(a) * r * 0.62, cy + math.sin(a) * r * 0.62)],
                    palette["ink_light"], 0.9, False))
    out.append(("text", cx, cy - r * 1.42, "N", int(r * 0.62), ink,
                "center", False, True, 0.0))
    return out


def _scale_primitives(gm: GameMap) -> List[tuple]:
    if not gm.scale_text.strip():
        return []
    palette = gm.palette()
    ink, paper = palette["ink"], palette["paper"]
    bar = min(gm.width * 0.16, 230.0)
    x0 = gm.width * 0.045
    y0 = gm.height - gm.height * 0.052
    height = max(7.0, gm.height * 0.011)
    out: List[tuple] = []
    segments = 4
    for i in range(segments):
        sx = x0 + bar / segments * i
        out.append(("polygon",
                    [(sx, y0), (sx + bar / segments, y0),
                     (sx + bar / segments, y0 + height), (sx, y0 + height)],
                    ink if i % 2 == 0 else paper, ink, 1.1, False))
    out.append(("text", x0 + bar / 2, y0 - height * 1.5, gm.scale_text,
                max(9, int(gm.height * 0.016)), ink, "center", True, False, 0.0))
    return out


def _border_primitives(gm: GameMap) -> List[tuple]:
    palette = gm.palette()
    ink, light = palette["ink"], palette["ink_light"]
    inset = max(10.0, min(gm.width, gm.height) * 0.018)
    out: List[tuple] = [
        ("polygon", [(inset, inset), (gm.width - inset, inset),
                     (gm.width - inset, gm.height - inset),
                     (inset, gm.height - inset)], "", ink, 2.4, False),
    ]
    inner = inset * 1.75
    out.append(("polygon", [(inner, inner), (gm.width - inner, inner),
                            (gm.width - inner, gm.height - inner),
                            (inner, gm.height - inner)], "", light, 1.0, False))
    return out


def _title_primitives(gm: GameMap) -> List[tuple]:
    if not gm.title_on_map or not gm.name.strip():
        return []
    palette = gm.palette()
    size = max(18, int(gm.height * 0.038))
    return [
        ("text", gm.width / 2.0, gm.height * 0.062, gm.name.upper(),
         size, palette["ink"], "center", False, True, size * 0.14),
    ]


# -- the main builder ------------------------------------------------------


def build_primitives(gm: GameMap, include_furniture: bool = True
                     ) -> List[tuple]:
    """Turn a map into an ordered display list. This is the single source of truth."""
    palette = gm.palette()
    ink = palette["ink"]
    light = palette["ink_light"]
    paper = palette["paper"]
    water = palette["water"]
    # Per-style terrain fills, so land stays clearly lighter than the sea in
    # every palette rather than only in the one it was tuned against.
    style_terrain: Dict[str, str] = palette.get("terrain", {}) or {}
    visible = gm.visible_layers()
    out: List[tuple] = []

    if gm.grid != "none":
        out.extend(_grid_primitives(gm))

    ordered = sorted(
        (s for s in gm.shapes if s.layer in visible),
        key=lambda s: PAINT_ORDER.get(s.kind, 50),
    )

    for shape in ordered:
        if len(shape.points) < 2:
            continue
        spec = TERRAIN.get(shape.kind, TERRAIN["land"])
        closed = shape.closed and spec["closed"] and len(shape.points) >= 3
        pts = _smooth(shape.points, 2 if len(shape.points) < 240 else 1)
        if closed and pts[0] != pts[-1]:
            pts = pts + [pts[0]]

        fill = shape.fill or style_terrain.get(shape.kind) or (spec["fill"] or "")
        outline = shape.outline or ink
        width = float(shape.width or 2.0)

        # Coastal halo: progressively wider, lighter strokes under the fill.
        # Scaled to the map so a large world map gets a proportionate glow.
        if spec.get("halo") and closed:
            band = max(4.0, min(9.0, gm.height * 0.006))
            halo_bands = list(reversed(palette["halo"]))
            for index, halo in enumerate(halo_bands):
                out.append(("line", pts, halo,
                            width + band * (len(halo_bands) - index), False))

        decor = spec.get("decor")
        if closed:
            out.append(("polygon", pts, fill, outline if fill else "",
                        width if fill else 0.0, False))
            if fill and shape.kind != "region":
                out.append(("line", pts, outline, width, False))
            elif shape.kind == "region":
                out.append(("line", pts, shape.outline or light,
                            max(1.2, width), True))
        else:
            if decor == "river":
                # Rivers need to read at a glance against a busy land fill, so
                # they get the deeper water tone and a little more weight.
                out.extend(_tapered_river(
                    pts, shape.fill or palette["water_deep"], width + 2.6
                ))
            elif decor == "road":
                out.append(("line", pts, shape.outline or light,
                            max(1.4, width), True))
            elif decor == "route":
                out.append(("line", pts, shape.outline or ink,
                            max(1.6, width), True))
                if len(pts) >= 2:
                    out.extend(_arrow_head(pts, shape.outline or ink,
                                           max(1.6, width)))
            elif decor == "wall":
                out.append(("line", pts, outline, max(2.4, width + 1.2), False))
                for (px, py), angle in _walk_path(pts, 14.0):
                    nx = math.cos(angle + math.pi / 2) * 3.4
                    ny = math.sin(angle + math.pi / 2) * 3.4
                    out.append(("line", [(px - nx, py - ny), (px + nx, py + ny)],
                                outline, 1.2, False))
            elif decor not in ("peaks", "hills"):
                out.append(("line", pts, outline, width, False))

        if decor == "peaks":
            out.extend(_peaks(pts, ink, paper, seed=gm.seed ^ (hash(shape.id) & 0xFF)))
        elif decor == "hills":
            out.extend(_hills(pts, ink))
        elif decor == "trees" and closed:
            out.extend(_trees(shape, ink, gm.seed))
        elif decor == "dots" and closed:
            out.extend(_dots(shape, light, gm.seed))
        elif decor == "squiggles" and closed:
            out.extend(_squiggles(shape, light, gm.seed))

        if shape.label.strip():
            cx, cy = shape.centroid()
            out.append(("text", cx, cy, shape.label,
                        max(11, int(gm.height * 0.019)), light, "center",
                        True, False, 2.0))

    accent = "#8a2f22" if gm.style in ("parchment", "treasure") else ink
    placement = layout_pin_labels(gm, visible) if gm.auto_place_labels else {}
    label_size = max(9, int(gm.height * 0.0155))
    for pin in gm.pins:
        if pin.layer not in visible:
            continue
        out.extend(pin_primitives(pin, ink, paper, accent))
        if not pin.label.strip():
            continue
        side, show = placement.get(pin.id, (pin.label_side or "e", True))
        if not show:
            continue
        lx, ly, anchor = pin_label_anchor(pin, side)
        out.append(("text", lx, ly, pin.label, label_size, ink, anchor,
                    False, pin.kind in ("capital", "city"), 0.0))

    for label in gm.labels:
        if label.layer not in visible:
            continue
        if not label.text.strip():
            continue
        out.append(("text", label.x, label.y, label.text, int(label.size),
                    label.color or ink, "center", label.italic, label.bold,
                    float(label.tracking)))

    if include_furniture:
        if gm.border:
            out.extend(_border_primitives(gm))
        if gm.compass:
            out.extend(_compass_primitives(gm))
        out.extend(_scale_primitives(gm))
        out.extend(_title_primitives(gm))

    return out


def _arrow_head(points: Sequence[Point], colour: str,
                width: float) -> List[tuple]:
    if len(points) < 2:
        return []
    (x0, y0), (x1, y1) = points[-2], points[-1]
    angle = math.atan2(y1 - y0, x1 - x0)
    size = max(7.0, width * 4.0)
    return [
        ("line", [(x1 - math.cos(angle - 0.42) * size,
                   y1 - math.sin(angle - 0.42) * size), (x1, y1)],
         colour, width, False),
        ("line", [(x1 - math.cos(angle + 0.42) * size,
                   y1 - math.sin(angle + 0.42) * size), (x1, y1)],
         colour, width, False),
    ]


# ==========================================================================
# PNG export (Pillow)
# ==========================================================================


def _font_folders() -> List[Path]:
    """
    Where to look for a real font file.

    Hardcoding C:/Windows/Fonts breaks on any machine where Windows is not on
    C:, and on the per-user font folder that installing a font without admin
    rights writes to. This is a portable tool, so it asks the system.
    """
    folders: List[Path] = []
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windir:
        folders.append(Path(windir) / "Fonts")
    local = os.environ.get("LOCALAPPDATA")
    if local:
        folders.append(Path(local) / "Microsoft" / "Windows" / "Fonts")
    # Not Windows: the usual places, so exported maps still get real type.
    folders += [Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
                Path.home() / ".fonts", Path("/Library/Fonts"),
                Path("/System/Library/Fonts")]
    return [f for f in folders if f.is_dir()]


def _find_font(italic: bool = False, bold: bool = False):
    from PIL import ImageFont

    if bold and italic:
        names = ["georgiaz.ttf", "timesbi.ttf", "constanz.ttf", "arialbi.ttf"]
    elif bold:
        names = ["georgiab.ttf", "timesbd.ttf", "constanb.ttf", "arialbd.ttf"]
    elif italic:
        names = ["georgiai.ttf", "timesi.ttf", "constani.ttf", "ariali.ttf"]
    else:
        names = ["georgia.ttf", "times.ttf", "constan.ttf", "arial.ttf",
                 "DejaVuSerif.ttf", "LiberationSerif-Regular.ttf"]
    for folder in _font_folders():
        for name in names:
            candidate = folder / name
            if candidate.exists():
                return str(candidate)
    return None


_FONT_CACHE: Dict[tuple, Any] = {}


def _font(size: int, italic: bool, bold: bool):
    from PIL import ImageFont

    key = (max(6, int(size)), bool(italic), bool(bold))
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    path = _find_font(italic, bold)
    try:
        font = ImageFont.truetype(path, key[0]) if path else ImageFont.load_default()
    except (OSError, ValueError):
        font = ImageFont.load_default()
    _FONT_CACHE[key] = font
    return font


def _parchment_background(gm: GameMap, scale: float):
    """Aged-paper texture: mottling plus a vignette. Deterministic per seed."""
    from PIL import Image, ImageDraw, ImageFilter

    palette = gm.palette()
    width = max(1, int(gm.width * scale))
    height = max(1, int(gm.height * scale))
    base = Image.new("RGB", (width, height), palette["paper"])

    if not palette.get("texture"):
        return base

    rng = random.Random(gm.seed or 7)
    blot = Image.new("L", (max(1, width // 6), max(1, height // 6)), 0)
    blot_draw = ImageDraw.Draw(blot)
    for _ in range(260):
        bx = rng.randint(0, blot.width)
        by = rng.randint(0, blot.height)
        br = rng.randint(2, max(3, blot.width // 14))
        blot_draw.ellipse([bx - br, by - br, bx + br, by + br],
                          fill=rng.randint(18, 78))
    blot = blot.filter(ImageFilter.GaussianBlur(radius=blot.width / 42.0))
    blot = blot.resize((width, height), Image.BILINEAR)
    dark = Image.new("RGB", (width, height), palette["paper_dark"])
    base = Image.composite(dark, base, blot.point(lambda v: min(255, v * 2)))

    if palette.get("vignette"):
        mask = Image.new("L", (width, height), 0)
        mask_draw = ImageDraw.Draw(mask)
        inset = int(min(width, height) * 0.06)
        mask_draw.rectangle([inset, inset, width - inset, height - inset],
                           fill=255)
        mask = mask.filter(
            ImageFilter.GaussianBlur(radius=max(6, min(width, height) * 0.05))
        )
        edge = Image.new("RGB", (width, height), palette["paper_dark"])
        base = Image.composite(base, edge, mask)

    return base


def render_png(gm: GameMap, path: Path | str, scale: float = 1.0,
               supersample: int = 2) -> Path:
    """
    Rasterise a map to PNG.

    Drawn at `supersample` times the requested size and downscaled, which is
    how the lines come out smooth - Pillow has no anti-aliased line drawing.
    """
    from PIL import Image, ImageDraw

    scale = max(0.2, min(4.0, float(scale)))
    ss = max(1, min(3, int(supersample)))
    effective = scale * ss

    image = _parchment_background(gm, effective)
    draw = ImageDraw.Draw(image, "RGBA")

    def sc(points: Iterable[Point]) -> List[Tuple[float, float]]:
        return [(p[0] * effective, p[1] * effective) for p in points]

    for prim in build_primitives(gm):
        head = prim[0]
        try:
            if head == "polygon":
                _kind, points, fill, outline, width, dash = prim
                pts = sc(points)
                if len(pts) < 3:
                    continue
                if fill:
                    draw.polygon(pts, fill=fill)
                if outline and width:
                    draw.line(pts + [pts[0]], fill=outline,
                              width=max(1, int(round(width * effective))),
                              joint="curve")
            elif head == "line":
                _kind, points, colour, width, dash = prim
                pts = sc(points)
                if len(pts) < 2 or not colour:
                    continue
                stroke = max(1, int(round(width * effective)))
                if dash:
                    for seg in _dash_segments(pts, 9.0 * effective,
                                              6.0 * effective):
                        if len(seg) >= 2:
                            draw.line(seg, fill=colour, width=stroke,
                                      joint="curve")
                else:
                    draw.line(pts, fill=colour, width=stroke, joint="curve")
            elif head == "ellipse":
                _kind, x0, y0, x1, y1, fill, outline, width = prim
                box = [x0 * effective, y0 * effective,
                       x1 * effective, y1 * effective]
                if box[2] < box[0]:
                    box[0], box[2] = box[2], box[0]
                if box[3] < box[1]:
                    box[1], box[3] = box[3], box[1]
                if box[2] - box[0] < 1 or box[3] - box[1] < 1:
                    continue
                draw.ellipse(box, fill=fill or None, outline=outline or None,
                             width=max(1, int(round(width * effective))))
            elif head == "text":
                _kind, x, y, text, size, colour, anchor, italic, bold, tracking = prim
                font = _font(int(size * effective), italic, bold)
                anchor_map = {"center": "mm", "w": "lm", "e": "rm",
                              "n": "ma", "s": "md"}
                if tracking and abs(tracking) > 0.4:
                    _draw_tracked(draw, x * effective, y * effective, text,
                                  font, colour, tracking * effective,
                                  anchor_map.get(anchor, "mm"))
                else:
                    draw.text((x * effective, y * effective), text, font=font,
                              fill=colour, anchor=anchor_map.get(anchor, "mm"))
        except (ValueError, TypeError, OSError):
            # One malformed primitive must not abandon the whole export.
            continue

    if ss > 1:
        target = (max(1, int(gm.width * scale)), max(1, int(gm.height * scale)))
        image = image.resize(target, Image.LANCZOS)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(str(path), "PNG", optimize=True)
    return path


def _draw_tracked(draw, x: float, y: float, text: str, font, colour: str,
                  tracking: float, anchor: str) -> None:
    """Letter-spaced text, which Pillow cannot do natively."""
    widths = [draw.textlength(ch, font=font) for ch in text]
    total = sum(widths) + tracking * max(0, len(text) - 1)
    if anchor.startswith("r"):
        cursor = x - total
    elif anchor.startswith("l"):
        cursor = x
    else:
        cursor = x - total / 2.0
    vertical = "m" if anchor.endswith("m") else anchor[-1]
    for ch, w in zip(text, widths):
        draw.text((cursor, y), ch, font=font, fill=colour,
                  anchor="l" + vertical)
        cursor += w + tracking


def _dash_segments(points: Sequence[Point], dash: float,
                   gap: float) -> List[List[Point]]:
    """Split a polyline into dashes. Pillow has no dash support."""
    dash = max(1.0, dash)
    gap = max(1.0, gap)
    segments: List[List[Point]] = []
    current: List[Point] = []
    drawing = True
    remaining = dash
    for a, b in zip(points, points[1:]):
        ax, ay = a
        bx, by = b
        seg = math.hypot(bx - ax, by - ay)
        if seg < 1e-9:
            continue
        travelled = 0.0
        guard = 0
        while travelled < seg:
            guard += 1
            if guard > 4000:
                break
            step = min(remaining, seg - travelled)
            t0 = travelled / seg
            t1 = (travelled + step) / seg
            p0 = (ax + (bx - ax) * t0, ay + (by - ay) * t0)
            p1 = (ax + (bx - ax) * t1, ay + (by - ay) * t1)
            if drawing:
                if not current:
                    current.append(p0)
                current.append(p1)
            travelled += step
            remaining -= step
            if remaining <= 1e-9:
                if drawing and current:
                    segments.append(current)
                    current = []
                drawing = not drawing
                remaining = dash if drawing else gap
    if current:
        segments.append(current)
    return segments


# ==========================================================================
# SVG export
# ==========================================================================


def _svg_escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def render_svg(gm: GameMap, path: Path | str) -> Path:
    palette = gm.palette()
    parts: List[str] = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{gm.width}" '
        f'height="{gm.height}" viewBox="0 0 {gm.width} {gm.height}">',
        f'<rect width="{gm.width}" height="{gm.height}" '
        f'fill="{palette["paper"]}"/>',
    ]

    def pts_attr(points: Sequence[Point]) -> str:
        return " ".join(f"{p[0]:.2f},{p[1]:.2f}" for p in points)

    for prim in build_primitives(gm):
        head = prim[0]
        if head == "polygon":
            _k, points, fill, outline, width, dash = prim
            if len(points) < 3:
                continue
            parts.append(
                f'<polygon points="{pts_attr(points)}" '
                f'fill="{fill or "none"}" stroke="{outline or "none"}" '
                f'stroke-width="{width:.2f}" stroke-linejoin="round"'
                + (' stroke-dasharray="9 6"' if dash else "") + "/>"
            )
        elif head == "line":
            _k, points, colour, width, dash = prim
            if len(points) < 2 or not colour:
                continue
            parts.append(
                f'<polyline points="{pts_attr(points)}" fill="none" '
                f'stroke="{colour}" stroke-width="{width:.2f}" '
                f'stroke-linecap="round" stroke-linejoin="round"'
                + (' stroke-dasharray="9 6"' if dash else "") + "/>"
            )
        elif head == "ellipse":
            _k, x0, y0, x1, y1, fill, outline, width = prim
            cx, cy = (x0 + x1) / 2.0, (y0 + y1) / 2.0
            rx, ry = abs(x1 - x0) / 2.0, abs(y1 - y0) / 2.0
            if rx < 0.4 or ry < 0.4:
                continue
            parts.append(
                f'<ellipse cx="{cx:.2f}" cy="{cy:.2f}" rx="{rx:.2f}" '
                f'ry="{ry:.2f}" fill="{fill or "none"}" '
                f'stroke="{outline or "none"}" stroke-width="{width:.2f}"/>'
            )
        elif head == "text":
            _k, x, y, text, size, colour, anchor, italic, bold, tracking = prim
            anchor_svg = {"center": "middle", "w": "start", "e": "end",
                          "n": "middle", "s": "middle"}.get(anchor, "middle")
            baseline = {"n": "hanging", "s": "auto"}.get(anchor, "central")
            parts.append(
                f'<text x="{x:.2f}" y="{y:.2f}" font-family="Georgia, serif" '
                f'font-size="{size}" fill="{colour}" '
                f'text-anchor="{anchor_svg}" dominant-baseline="{baseline}"'
                + (' font-style="italic"' if italic else "")
                + (' font-weight="bold"' if bold else "")
                + (f' letter-spacing="{tracking:.2f}"' if tracking else "")
                + f">{_svg_escape(text)}</text>"
            )

    parts.append("</svg>")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


# ==========================================================================
# Word export
# ==========================================================================


def render_docx(gm: GameMap, png_path: Path, out_path: Path,
                location_lookup: Optional[Dict[str, str]] = None) -> Path:
    """
    A Word document holding the map image plus a legend of every pin.

    Keeps maps inside the "everything is a Word document" promise, and gives
    you something printable to keep beside you while writing.
    """
    from docx.shared import Inches

    from .atomic import save_via_atomic
    from . import docxio

    doc = docxio.sheet_document(margin=0.7)
    docxio.add_title_block(
        doc, gm.name, f"{MAP_KINDS.get(gm.kind, gm.kind)} - "
                      f"{STYLES.get(gm.style, {}).get('label', gm.style)}",
        "MAP",
    )

    if png_path.exists():
        # 7.1in fits US Letter with 0.7in margins.
        doc.add_picture(str(png_path), width=Inches(7.1))

    if gm.scale_text.strip():
        docxio.add_note(doc, f"Scale: {gm.scale_text}")

    if gm.pins:
        doc.add_heading("Legend", level=2)
        rows = []
        for pin in sorted(gm.pins, key=lambda p: (p.kind, p.label.lower())):
            linked = ""
            if pin.entity_id and location_lookup:
                linked = location_lookup.get(pin.entity_id, "")
            rows.append([
                pin.label or "(unnamed)",
                PIN_KINDS.get(pin.kind, pin.kind),
                linked,
                pin.notes,
            ])
        docxio.add_data_table(
            doc, ["Place", "Type", "Location sheet", "Notes"], rows,
            widths=[1.7, 1.2, 1.5, 2.6],
        )

    regions = [s for s in gm.shapes if s.label.strip()]
    if regions:
        doc.add_heading("Named areas", level=2)
        docxio.add_data_table(
            doc, ["Name", "Kind"],
            [[s.label, TERRAIN.get(s.kind, {}).get("label", s.kind)]
             for s in regions],
            widths=[3.0, 2.0],
        )

    if gm.notes.strip():
        doc.add_heading("Notes", level=2)
        for para in gm.notes.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    counts: Dict[str, int] = {}
    for shape in gm.shapes:
        counts[shape.kind] = counts.get(shape.kind, 0) + 1
    if counts:
        doc.add_heading("Contents", level=2)
        summary = ", ".join(
            f"{n} x {TERRAIN.get(k, {}).get('label', k).lower()}"
            for k, n in sorted(counts.items())
        )
        doc.add_paragraph(f"{summary}. {len(gm.pins)} pins, "
                          f"{len(gm.labels)} labels.")

    return save_via_atomic(out_path, doc.save)


# ==========================================================================
# Persistence
# ==========================================================================

MAP_SUFFIX = ".map.json"


def map_path(folder: Path, gm: GameMap) -> Path:
    return Path(folder) / f"{safe_filename(gm.name, 'Map')}{MAP_SUFFIX}"


def save_map(folder: Path, gm: GameMap) -> Path:
    gm.touch()
    path = map_path(folder, gm)
    write_json_atomic(path, gm.to_json())
    return path


def load_map(path: Path) -> Optional[GameMap]:
    data = read_json(path)
    if not isinstance(data, dict):
        return None
    try:
        return GameMap.from_json(data)
    except (TypeError, ValueError):
        return None


def list_maps(folder: Path) -> List[Tuple[str, Path]]:
    folder = Path(folder)
    if not folder.is_dir():
        return []
    out: List[Tuple[str, Path]] = []
    for path in sorted(folder.glob(f"*{MAP_SUFFIX}")):
        data = read_json(path, {}) or {}
        out.append((data.get("name") or path.name[:-len(MAP_SUFFIX)], path))
    return out


def delete_map(path: Path) -> bool:
    try:
        Path(path).unlink(missing_ok=True)
        return True
    except OSError:
        return False


# ==========================================================================
# Starter content
# ==========================================================================


def starter_map(name: str = "The Known World", kind: str = "world",
                style: str = "parchment") -> GameMap:
    """
    A new map with a plausible coastline already on it.

    An empty canvas is paralysing; something to push around is not. Generated
    from a fixed seed so it is the same every time and can be deleted in one
    click if unwanted.
    """
    gm = GameMap(name=name, kind=kind, style=style)
    if kind in ("city", "building", "dungeon", "battle"):
        gm.grid = "square"
        gm.grid_size = 60
        gm.compass = kind != "building"
        return gm

    rng = random.Random(11)
    cx, cy = gm.width * 0.47, gm.height * 0.52
    rx, ry = gm.width * 0.31, gm.height * 0.33
    coast: List[Point] = []
    steps = 46
    for i in range(steps):
        a = 2 * math.pi * i / steps
        wobble = 1.0 + 0.22 * math.sin(a * 3.0 + 0.7) + rng.uniform(-0.09, 0.09)
        coast.append((cx + math.cos(a) * rx * wobble,
                      cy + math.sin(a) * ry * wobble))
    gm.shapes.append(Shape(kind="land", points=coast, closed=True,
                           label="", layer="Base"))

    gm.shapes.append(Shape(
        kind="mountains", closed=False, layer="Base",
        points=[(cx - rx * 0.45, cy - ry * 0.42), (cx - rx * 0.1, cy - ry * 0.52),
                (cx + rx * 0.25, cy - ry * 0.38), (cx + rx * 0.5, cy - ry * 0.16)],
    ))
    gm.shapes.append(Shape(
        kind="forest", closed=True, layer="Base",
        points=[(cx - rx * 0.62, cy + ry * 0.05), (cx - rx * 0.2, cy - ry * 0.05),
                (cx - rx * 0.08, cy + ry * 0.38), (cx - rx * 0.55, cy + ry * 0.46)],
    ))
    gm.shapes.append(Shape(
        kind="river", closed=False, layer="Base",
        points=[(cx - rx * 0.05, cy - ry * 0.44), (cx + rx * 0.04, cy - ry * 0.1),
                (cx - rx * 0.06, cy + ry * 0.22), (cx + rx * 0.12, cy + ry * 0.62),
                (cx + rx * 0.3, cy + ry * 0.95)],
    ))
    gm.pins.append(Pin(x=cx + rx * 0.14, y=cy + ry * 0.6, kind="capital",
                       label="Capital", label_side="e", size=8))
    gm.pins.append(Pin(x=cx - rx * 0.5, y=cy + ry * 0.3, kind="town",
                       label="A town", label_side="w"))
    gm.pins.append(Pin(x=cx + rx * 0.55, y=cy - ry * 0.05, kind="ruin",
                       label="Ruins", label_side="e"))
    gm.labels.append(MapLabel(x=cx, y=cy + ry * 1.28, text="THE OCEAN",
                              size=20, italic=True, tracking=6.0))
    gm.scale_text = "100 leagues"
    return gm


NAME_SYLLABLES = {
    "northern": (["Bran", "Thor", "Sker", "Vald", "Hald", "Grim", "Fjor",
                  "Ulf", "Stein", "Orm", "Rag", "Sig"],
                 ["dal", "vik", "holm", "gard", "fell", "borg", "mark",
                  "stad", "ness", "fjord", "heim", "by"]),
    "southern": (["Cala", "Meri", "Sera", "Vala", "Alma", "Tara", "Sola",
                  "Bela", "Cora", "Lira", "Nera", "Mira"],
                 ["nova", "mar", "vera", "londe", "riva", "sette", "monte",
                  "corte", "valle", "porto", "bella", "sana"]),
    "elvish": (["Ael", "Cel", "Ely", "Fin", "Gal", "Ith", "Lor", "Mith",
                "Nim", "Sil", "Thal", "Vae"],
               ["andor", "ariel", "wen", "rond", "aloth", "ithil", "dor",
                "las", "riel", "orn", "eth", "amar"]),
    "harsh": (["Kra", "Zor", "Gnash", "Vrak", "Mog", "Skul", "Thrag",
               "Ur", "Bok", "Drez", "Hak", "Nur"],
              ["gul", "zar", "dun", "rok", "grim", "mor", "kath", "thul",
               "gash", "var", "nak", "dread"]),
    "plain": (["Ash", "Black", "Cold", "Green", "Grey", "High", "Long",
               "North", "Oak", "Salt", "Stone", "White", "Red", "Deep"],
              ["ford", "bridge", "field", "wood", "hill", "water", "gate",
               "haven", "moor", "brook", "ridge", "reach", "hollow", "cross"]),
}


#: The shipped lists, kept so the editable file can always be restored.
DEFAULT_NAME_SYLLABLES = {
    key: (list(start), list(end))
    for key, (start, end) in NAME_SYLLABLES.items()
}


def name_styles_file(folder: Path | str) -> Path:
    return Path(folder) / "Name Styles.json"


def load_name_styles(folder: Path | str) -> Path:
    """
    Read the writer's own name lists, creating the starter file if absent.

    The generator used to be five hardcoded lists, which is fine until your
    world does not sound like any of them. This is a plain JSON file in the
    Maps folder: add a style, delete one, replace every syllable with your own.
    Anything malformed is ignored in favour of the shipped lists rather than
    breaking name generation.
    """
    path = name_styles_file(folder)
    if not path.exists():
        try:
            write_json_atomic(path, {
                "_note": "Each style has a list of beginnings and a list of "
                         "endings. A name is one of each, joined. Add your own "
                         "styles, or replace these entirely.",
                **{key: {"start": start, "end": end}
                   for key, (start, end) in DEFAULT_NAME_SYLLABLES.items()},
            })
        except OSError:
            return path

    data = read_json(path, None)
    if not isinstance(data, dict):
        return path
    merged: Dict[str, Tuple[List[str], List[str]]] = {}
    for key, value in data.items():
        if key.startswith("_") or not isinstance(value, dict):
            continue
        start = value.get("start")
        end = value.get("end")
        if isinstance(start, list) and isinstance(end, list) and start and end:
            merged[str(key)] = ([str(s) for s in start if str(s).strip()],
                                [str(e) for e in end if str(e).strip()])
    if merged:
        NAME_SYLLABLES.clear()
        NAME_SYLLABLES.update(merged)
    return path


def name_styles() -> List[str]:
    return sorted(NAME_SYLLABLES)


def generate_name(flavour: str = "plain", seed: Optional[int] = None) -> str:
    """A place-name generator. Trivial, and saves a surprising amount of time."""
    if flavour not in NAME_SYLLABLES and NAME_SYLLABLES:
        flavour = next(iter(NAME_SYLLABLES))
    first, second = NAME_SYLLABLES.get(
        flavour, next(iter(NAME_SYLLABLES.values())))
    rng = random.Random(seed)
    name = rng.choice(first) + rng.choice(second)
    if flavour == "plain" and rng.random() < 0.25 and "plain" in NAME_SYLLABLES:
        name += " " + rng.choice(["Keep", "Cross", "End", "Mill", "Watch"])
    return name


def generate_names(flavour: str = "plain", count: int = 12) -> List[str]:
    out: List[str] = []
    guard = 0
    while len(out) < count and guard < count * 20:
        guard += 1
        candidate = generate_name(flavour)
        if candidate not in out:
            out.append(candidate)
    return out
