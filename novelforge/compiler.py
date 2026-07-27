"""
Compile: every scene document assembled into one manuscript.

Output is standard manuscript format, the thing agents and editors expect:
US Letter, one-inch margins, 12pt, double spaced, half-inch first-line indent,
a title page carrying contact details and word count, a running header of
`Surname / TITLE / page`, chapters starting on fresh pages, and `#` between
scenes.

Run-level italics and bold survive the trip; everything else is normalised to
the manuscript style, which is the point of compiling.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from .atomic import safe_filename, save_via_atomic
from .config import settings
from . import docxio
from .model import Scene

# --------------------------------------------------------------------------
# Number words for "Chapter Seventeen"
# --------------------------------------------------------------------------

_ONES = [
    "Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight",
    "Nine", "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen",
    "Sixteen", "Seventeen", "Eighteen", "Nineteen",
]
_TENS = [
    "", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy",
    "Eighty", "Ninety",
]


def number_word(n: int) -> str:
    """1 -> 'One', 42 -> 'Forty-Two'. Falls back to digits above 999."""
    if n < 0:
        return str(n)
    if n < 20:
        return _ONES[n]
    if n < 100:
        tens, ones = divmod(n, 10)
        return _TENS[tens] + (f"-{_ONES[ones]}" if ones else "")
    if n < 1000:
        hundreds, rest = divmod(n, 100)
        out = f"{_ONES[hundreds]} Hundred"
        return f"{out} {number_word(rest)}" if rest else out
    return str(n)


# --------------------------------------------------------------------------
# Options
# --------------------------------------------------------------------------


@dataclass
class CompileOptions:
    title_page: bool = True
    table_of_contents: bool = False
    chapter_headings: bool = True
    chapter_numbering: str = "word"      # word | digit | none
    include_chapter_titles: bool = True
    scene_separator: str = "#"
    include_synopses: bool = False       # working-draft mode
    include_status_notes: bool = False
    running_header: bool = True
    the_end: bool = True
    only_chapter_ids: Optional[Sequence[str]] = None
    font: str = ""
    font_size: int = 0
    line_spacing: float = 0.0
    margin: float = 0.0
    first_line_indent: float = -1.0
    contact_block: str = ""              # address / phone / email, one per line

    def resolved(self) -> "CompileOptions":
        """Fill blanks from user settings."""
        return CompileOptions(
            title_page=self.title_page,
            table_of_contents=self.table_of_contents,
            chapter_headings=self.chapter_headings,
            chapter_numbering=self.chapter_numbering,
            include_chapter_titles=self.include_chapter_titles,
            scene_separator=self.scene_separator or settings["scene_separator"],
            include_synopses=self.include_synopses,
            include_status_notes=self.include_status_notes,
            running_header=self.running_header,
            the_end=self.the_end,
            only_chapter_ids=self.only_chapter_ids,
            font=self.font or settings["manuscript_font"],
            font_size=self.font_size or settings["manuscript_font_size"],
            line_spacing=self.line_spacing or settings["manuscript_line_spacing"],
            margin=self.margin or settings["manuscript_margin"],
            first_line_indent=(
                self.first_line_indent
                if self.first_line_indent >= 0
                else settings["manuscript_first_line_indent"]
            ),
            contact_block=self.contact_block,
        )


@dataclass
class CompileResult:
    path: Path
    words: int
    chapters: int
    scenes: int
    skipped: List[str]
    # Scenes whose .docx could not be read. Silence here would mean a scene
    # vanishing from the manuscript with no indication at all.
    missing: List[str] = field(default_factory=list)

    @property
    def summary(self) -> str:
        text = (
            f"{self.words:,} words - {self.chapters} chapters, "
            f"{self.scenes} scenes"
        )
        if self.skipped:
            text += f" ({len(self.skipped)} skipped)"
        if self.missing:
            text += f" - WARNING: {len(self.missing)} unreadable"
        return text


# --------------------------------------------------------------------------
# Compile
# --------------------------------------------------------------------------


def compile_manuscript(project, options: Optional[CompileOptions] = None,
                       destination: Optional[Path] = None) -> CompileResult:
    """Assemble the whole book into a single .docx."""
    opts = (options or CompileOptions()).resolved()
    data = project.data

    doc = docxio.manuscript_document(
        font=opts.font,
        size=opts.font_size,
        line_spacing=opts.line_spacing,
        margin=opts.margin,
        first_line_indent=opts.first_line_indent,
    )
    # Remove the empty paragraph python-docx starts a Document with, so the
    # title page begins exactly where we put it.
    body = doc.element.body
    for para in list(doc.paragraphs):
        body.remove(para._p)

    section = doc.sections[0]
    chapters = [
        c for c in data.ordered_chapters()
        if c.include_in_compile
        and (opts.only_chapter_ids is None or c.id in opts.only_chapter_ids)
    ]
    included_ids = {c.id for c in chapters}
    scenes = [
        s for s in data.ordered_scenes()
        if s.include_in_compile and s.chapter_id in included_ids
    ]
    total_words = sum(s.word_count for s in scenes)

    skipped = [
        s.title for s in data.ordered_scenes()
        if not s.include_in_compile or s.chapter_id not in included_ids
    ]

    # -- running header ------------------------------------------------
    if opts.running_header:
        section.different_first_page_header_footer = bool(opts.title_page)
        _write_running_header(section, data, opts)
        if opts.title_page:
            # Blank the first-page header so the title page stays clean.
            first = section.first_page_header
            for para in first.paragraphs:
                para.text = ""

    # -- title page ----------------------------------------------------
    if opts.title_page:
        _write_title_page(doc, data, total_words, opts)

    # -- table of contents ---------------------------------------------
    if opts.table_of_contents:
        _write_toc(doc, opts, page_break=opts.title_page)

    # -- body ----------------------------------------------------------
    #
    # Chapters with nothing in them are dropped BEFORE numbering, so an empty
    # chapter cannot consume "Chapter Four" and leave a gap in the sequence.
    populated = [
        (chapter, [s for s in data.scenes_in(chapter.id) if s.include_in_compile])
        for chapter in chapters
    ]
    populated = [(c, s) for c, s in populated if s]

    first_chapter = True
    scene_count = 0
    missing: List[str] = []
    for index, (chapter, chapter_scenes) in enumerate(populated, start=1):
        if not first_chapter or opts.title_page or opts.table_of_contents:
            _page_break(doc)
        first_chapter = False

        if opts.chapter_headings:
            _write_chapter_heading(doc, chapter, index, opts)

        for scene_index, scene in enumerate(chapter_scenes):
            if scene_index:
                _write_separator(doc, opts)
            if opts.include_synopses and scene.synopsis:
                _write_working_note(doc, f"[{scene.title}] {scene.synopsis}", opts)
            elif opts.include_status_notes:
                _write_working_note(
                    doc, f"[{scene.title} - {scene.status}]", opts
                )
            if not _write_scene(doc, project, scene, opts):
                missing.append(scene.title)
            scene_count += 1

    # -- the end -------------------------------------------------------
    if opts.the_end and scene_count:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.space_before = Pt(opts.font_size * 2)
        para.add_run("THE END")

    if opts.table_of_contents:
        docxio.request_field_update(doc)

    target = Path(destination) if destination else project.compiled_path
    save_via_atomic(target, doc.save)

    return CompileResult(
        path=target,
        words=total_words,
        chapters=len(populated),
        scenes=scene_count,
        skipped=skipped,
        missing=missing,
    )


# --------------------------------------------------------------------------
# Pieces
# --------------------------------------------------------------------------


def _page_break(doc) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    para.add_run().add_break(WD_BREAK.PAGE)


def _write_running_header(section, data, opts: CompileOptions) -> None:
    """`Surname / TITLE / 4`, right-aligned, on every page but the first."""
    header = section.header
    header.is_linked_to_previous = False
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.line_spacing = 1.0

    surname = data.author_surname or (
        data.author.strip().split()[-1] if data.author.strip() else "Author"
    )
    keyword = _title_keyword(data.title)
    run = para.add_run(f"{surname} / {keyword} / ")
    run.font.name = opts.font
    run.font.size = Pt(opts.font_size)
    docxio.add_field(para, "PAGE", "1")
    for run in para.runs:
        run.font.name = opts.font
        run.font.size = Pt(opts.font_size)


def _title_keyword(title: str) -> str:
    """One memorable word from the title, upper-cased, for the header."""
    skip = {"a", "an", "the", "of", "and", "or", "in", "on", "to", "for"}
    words = [w.strip(".,:;!?'\"") for w in (title or "").split()]
    words = [w for w in words if w and w.lower() not in skip]
    return (words[0] if words else (title or "UNTITLED")).upper()


def _write_title_page(doc, data, total_words: int, opts: CompileOptions) -> None:
    """Contact block top-left, word count top-right, title a third of the way down."""
    contact_lines = [line for line in (opts.contact_block or "").splitlines() if line.strip()]
    if not contact_lines:
        contact_lines = [data.author or "Author name", "Address", "Email", "Phone"]

    # Word count sits on the first line, right-aligned.
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.space_after = Pt(0)
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.add_run(f"About {_round_words(total_words):,} words")

    for line in contact_lines:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(0)
        para.add_run(line)

    # Push the title down the page.
    for _ in range(9):
        blank = doc.add_paragraph()
        blank.paragraph_format.first_line_indent = Pt(0)
        blank.paragraph_format.space_after = Pt(0)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.first_line_indent = Pt(0)
    title_para.add_run(data.title.upper())

    if data.subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.first_line_indent = Pt(0)
        sub.add_run(data.subtitle)

    blank = doc.add_paragraph()
    blank.paragraph_format.first_line_indent = Pt(0)

    byline = doc.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.paragraph_format.first_line_indent = Pt(0)
    byline.add_run(f"by {data.author or 'Author'}")

    if data.series:
        series = doc.add_paragraph()
        series.alignment = WD_ALIGN_PARAGRAPH.CENTER
        series.paragraph_format.first_line_indent = Pt(0)
        label = f"{data.series}"
        if data.book_number:
            label += f", Book {data.book_number}"
        series.add_run(label)


def _round_words(n: int) -> int:
    """Manuscript word counts are given rounded - to 1,000 above 10k."""
    if n >= 10000:
        return int(round(n / 1000.0)) * 1000
    if n >= 1000:
        return int(round(n / 500.0)) * 500
    return n


def _write_toc(doc, opts: CompileOptions, page_break: bool = True) -> None:
    # Only break if something precedes us, or the contents page is preceded by
    # a blank one.
    if page_break:
        _page_break(doc)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Pt(0)
    heading.add_run("CONTENTS")
    doc.add_paragraph().paragraph_format.first_line_indent = Pt(0)

    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    docxio.add_field(
        para,
        r'TOC \o "1-2" \h \z \u',
        "Right-click and choose Update Field to build the contents.",
    )


def _write_chapter_heading(doc, chapter, index: int, opts: CompileOptions) -> None:
    """Chapter heading, a third of the way down a fresh page."""
    for _ in range(4):
        blank = doc.add_paragraph()
        blank.paragraph_format.first_line_indent = Pt(0)
        blank.paragraph_format.space_after = Pt(0)

    label = ""
    if opts.chapter_numbering == "word":
        label = f"Chapter {number_word(index)}"
    elif opts.chapter_numbering == "digit":
        label = f"Chapter {index}"

    if label:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.keep_with_next = True
        # Heading 1 style is what a TOC field picks up.
        para.style = doc.styles["Heading 1"]
        run = para.add_run(label)
        run.font.name = opts.font
        run.font.size = Pt(opts.font_size)
        run.font.bold = False
        run.font.color.rgb = docxio.INK

    title = (chapter.title or "").strip()
    generic = title.lower().startswith("chapter") or not title
    if opts.include_chapter_titles and title and not generic:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.keep_with_next = True
        if not label:
            para.style = doc.styles["Heading 1"]
            run = para.add_run(title)
            run.font.name = opts.font
            run.font.size = Pt(opts.font_size)
            run.font.bold = False
            run.font.color.rgb = docxio.INK
        else:
            para.add_run(title)

    blank = doc.add_paragraph()
    blank.paragraph_format.first_line_indent = Pt(0)
    blank.paragraph_format.space_after = Pt(0)


def _write_separator(doc, opts: CompileOptions) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Pt(0)
    para.add_run(opts.scene_separator)


def _write_working_note(doc, text: str, opts: CompileOptions) -> None:
    """A bracketed editorial note, for draft-only compiles."""
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(text)
    run.italic = True
    run.font.size = Pt(max(9, opts.font_size - 2))
    run.font.color.rgb = docxio.MUTED


def _write_scene(doc, project, scene: Scene, opts: CompileOptions) -> bool:
    """
    Copy a scene's prose in, preserving italics and bold.

    Returns False if the document is missing or unreadable, so the caller can
    warn instead of quietly producing a manuscript with a hole in it.
    """
    if not scene.docx:
        return not scene.word_count
    path = project.abs(scene.docx)
    if not path.exists():
        return False
    blocks = docxio.read_paragraphs_with_format(path)
    if not blocks and scene.word_count:
        # The manifest says there are words but nothing came back.
        return False
    for block in blocks:
        para = doc.add_paragraph()
        if block.get("centred") or (
            len(block["runs"]) == 1 and block["runs"][0]["text"].strip() == "#"
        ):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
        docxio.clone_runs(block, para)
        for run in para.runs:
            run.font.name = opts.font
            run.font.size = Pt(opts.font_size)
    return True


# --------------------------------------------------------------------------
# Plain text export - useful for word processors, diffing and dictation apps
# --------------------------------------------------------------------------


def export_plain_text(project, destination: Optional[Path] = None) -> Tuple[Path, int]:
    from .atomic import write_text_atomic

    data = project.data
    lines: List[str] = [data.title.upper(), ""]
    if data.author:
        lines += [f"by {data.author}", ""]
    lines.append("")

    words = 0
    # Filter first, then number - same reason as the docx compile: an excluded
    # or empty chapter must not consume a chapter number.
    populated = [
        (c, [s for s in data.scenes_in(c.id) if s.include_in_compile])
        for c in data.ordered_chapters() if c.include_in_compile
    ]
    populated = [(c, s) for c, s in populated if s]
    for index, (chapter, chapter_scenes) in enumerate(populated, start=1):
        lines += ["", f"CHAPTER {number_word(index).upper()}", ""]
        if chapter.title and not chapter.title.lower().startswith("chapter"):
            lines += [chapter.title, ""]
        for i, scene in enumerate(chapter_scenes):
            if i:
                lines += ["", "#", ""]
            text = docxio.read_prose(project.abs(scene.docx))
            words += docxio.word_count(text)
            lines.append(text)
    lines += ["", "THE END", ""]

    target = Path(destination) if destination else (
        project.folder("compiled")
        / f"{safe_filename(data.title, 'Manuscript')} - Manuscript.txt"
    )
    write_text_atomic(target, "\n".join(lines))
    return target, words


def compile_outline_only(project, destination: Optional[Path] = None) -> Path:
    """
    A synopsis-only document: every scene's synopsis in reading order.

    Reads like a treatment, and is the fastest way to check structure.
    """
    data = project.data
    doc = docxio.sheet_document()
    docxio.add_title_block(
        doc, data.title, "Scene-by-scene treatment", "OUTLINE",
    )
    for index, chapter in enumerate(data.ordered_chapters(), start=1):
        scenes = data.scenes_in(chapter.id)
        if not scenes:
            continue
        doc.add_heading(f"{index}. {chapter.title}", level=2)
        if chapter.synopsis:
            docxio.add_note(doc, chapter.synopsis)
        for scene in scenes:
            doc.add_heading(scene.title, level=3)
            pov = data.entity(scene.pov_id)
            meta = " | ".join(
                bit for bit in [
                    f"POV: {pov.name}" if pov else "",
                    scene.status,
                    f"{scene.word_count:,} words",
                    scene.story_date,
                ] if bit
            )
            if meta:
                docxio.add_note(doc, meta)
            doc.add_paragraph(scene.synopsis or "(no synopsis yet)")
            for label, value in (
                ("Goal", scene.goal), ("Conflict", scene.conflict),
                ("Disaster", scene.disaster), ("Reaction", scene.reaction),
                ("Dilemma", scene.dilemma), ("Decision", scene.decision),
            ):
                if value:
                    para = doc.add_paragraph()
                    para.add_run(f"{label}: ").bold = True
                    para.add_run(value)

    target = Path(destination) if destination else (
        project.folder("compiled")
        / f"{safe_filename(data.title, 'Novel')} - Treatment.docx"
    )
    return save_via_atomic(target, doc.save)
